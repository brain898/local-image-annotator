"""Real >40 MiB PPTX upload/render/download in isolated Chrome; no model call."""
from pathlib import Path
import sys,tempfile,threading,io,json
sys.path.insert(0,str(Path(__file__).resolve().parents[1]))
import numpy as np
from PIL import Image
from pptx import Presentation
from pptx.util import Inches
from playwright.sync_api import sync_playwright
import core,server
from ppt_package import Package

def run():
    old=core.DATA;report={'at':core.now(),'isolated':True,'model_called':False}
    with tempfile.TemporaryDirectory(prefix='ppt-large-browser-') as tmp:
        tmp=Path(tmp);core.DATA=tmp/'data';core.init()
        noise=np.random.default_rng(42).integers(0,256,(4096,4096,3),dtype=np.uint8);buf=io.BytesIO();Image.fromarray(noise).save(buf,format='PNG',compress_level=0);del noise
        p=Presentation();s=p.slides.add_slide(p.slide_layouts[6]);s.shapes.add_picture(io.BytesIO(buf.getvalue()),0,0,width=p.slide_width,height=p.slide_height)
        shape=s.shapes.add_textbox(Inches(.5),Inches(.5),Inches(5),Inches(.7));shape.text='Large upload 2025'
        sample=tmp/'大文件验证.pptx';p.save(sample);del p,buf
        report['input_bytes']=sample.stat().st_size;assert report['input_bytes']>40*1024*1024
        instance=server.LocalServer(('127.0.0.1',0),server.Handler);threading.Thread(target=instance.serve_forever,daemon=True).start()
        try:
            with sync_playwright() as pw:
                browser=pw.chromium.launch(channel='chrome',headless=True);page=browser.new_page(viewport={'width':1480,'height':1000});errors=[];requests=[];page.on('pageerror',lambda e:errors.append(str(e)))
                page.on('request',lambda r:requests.append({'url':r.url,'content_type':r.headers.get('content-type')}) if '/api/ppt/upload' in r.url else None)
                page.goto(f'http://127.0.0.1:{instance.server_port}/ppt.html');page.locator('#file').set_input_files(str(sample))
                page.wait_for_function("() => P.data.decks[0]?.status==='ready' || P.data.decks[0]?.status==='failed'",timeout=180000)
                d=page.evaluate('P.data.decks[0]');assert d['status']=='ready',d.get('error');assert page.locator('.thumb').count()==1;assert page.locator('#page-image').evaluate('(e)=>e.naturalWidth')==1600
                assert len(requests)==1 and requests[0]['url'].endswith('/api/ppt/upload-file');assert requests[0]['content_type']=='application/octet-stream'
                original=Path(d['original_path']);assert core.sha(original)==core.sha(sample)
                page.locator('#result').evaluate('(e)=>e.open=true')
                with page.expect_download() as download:page.get_by_role('button',name='下载原文件',exact=True).click()
                target=tmp/'download.pptx';download.value.save_as(str(target));assert core.sha(target)==core.sha(sample)
                pack=Package(target);assert pack.object(1,'3')['text']=='Large upload 2025';assert not pack.parts.edits
                report.update(binary_request=True,original_hash_preserved=True,powerpoint_real_render=True,download_hash_preserved=True,lazy_parts=True,page_errors=errors);assert not errors;browser.close()
        finally:instance.shutdown();instance.server_close();core.DATA=old
    core.dump(core.ROOT/'evidence/ppt-large-upload.json',report);print(json.dumps(report,ensure_ascii=False,indent=2))
if __name__=='__main__':run()
