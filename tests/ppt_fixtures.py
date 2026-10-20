"""Isolated synthetic decks. Never presented as real image-model evidence."""
from pathlib import Path
import io
from PIL import Image, ImageDraw, ImageFont

def make_deck(path):
    from pptx import Presentation
    from pptx.util import Inches, Pt
    from pptx.dml.color import RGBColor
    p=Presentation(); p.slide_width=Inches(12); p.slide_height=Inches(6.75)
    def slide(): return p.slides.add_slide(p.slide_layouts[6])
    def textbox(s,text,x=1,y=1,w=6,h=1):
        shape=s.shapes.add_textbox(Inches(x),Inches(y),Inches(w),Inches(h)); shape.text=text
        for para in shape.text_frame.paragraphs:
            for run in para.runs: run.font.size=Pt(28)
        return shape
    s=slide(); shape=textbox(s,'')
    for text,bold in [('Plan ',False),('20',True),('25',False)]:
        r=shape.text_frame.paragraphs[0].add_run(); r.text=text; r.font.size=Pt(32); r.font.bold=bold; r.font.color.rgb=RGBColor.from_string('24483F')
    s.notes_slide.notes_text_frame.text='fixture notes, preserve as data'
    textbox(s,'Unchanged object',1,4)
    s=slide(); table=s.shapes.add_table(2,2,Inches(1),Inches(1),Inches(8),Inches(2)).table
    for row,values in enumerate([['Year','Value'],['2025','123']]):
        for col,value in enumerate(values): table.cell(row,col).text=value
    im=Image.new('RGB',(1280,720),'#e2e9e2'); d=ImageDraw.Draw(im)
    import random
    r=random.Random(100)
    for i in range(700):
        x=r.randrange(1280); y=r.randrange(720); color=tuple(r.randrange(60,190) for _ in range(3))
        d.rectangle((x,y,x+r.randrange(2,9),y+r.randrange(2,9)),fill=color)
    d.rectangle((420,260,830,430),fill='white')
    font=ImageFont.truetype('C:/Windows/Fonts/arial.ttf',56); d.text((460,310),'2025',font=font,fill='black')
    raw=io.BytesIO(); im.save(raw,format='PNG'); data=raw.getvalue()
    s=slide(); s.shapes.add_picture(io.BytesIO(data),0,0,width=p.slide_width,height=p.slide_height)
    s=slide(); s.shapes.add_picture(io.BytesIO(data),0,0,width=p.slide_width,height=p.slide_height); textbox(s,'Overlay 2025',.5,.3,5,1)
    s=slide(); pic=s.shapes.add_picture(io.BytesIO(data),Inches(3),Inches(2),width=Inches(6),height=Inches(3.375)); pic.crop_left=.1; pic.crop_top=.1; pic.crop_right=.1; pic.crop_bottom=.1; pic.rotation=17
    s=slide(); textbox(s,'2025 and 2025')
    s=slide(); textbox(s,'Keep page unchanged')
    p.save(path); return Path(path)


def add_mock_analysis(manifest):
    """Program test stub, NEVER evidence that a model viewed the slide."""
    import json, core, ppt_workflow as ppt, ppt_analysis
    from ppt_package import Package
    j=core.read('ppt_job',manifest['job_id'])
    result=ppt.context(j['id'],manifest['lease'],include_ocr=False)
    ctx=json.loads(Path(result['context_path']).read_text(encoding='utf-8'))
    manifest['context_sha256']=result['context_sha256']
    manifest['page_reviews']=[dict(page=p['number'],preview_sha256=p['visual_reference']['sha256'],
        observation='程序测试桩：未由模型查看；隔离样例视觉另行核对') for p in ctx['pages']]
    pack=Package(j['base_path'])
    for item in manifest['items']:
        if item['method']=='confirm': continue
        clean=ppt.validate_item(pack,j,item); obj=pack.object(item['page'],item['object_id'])
        objects=pack.inspect()['pages'][item['page']-1]['objects']
        item['target_review']=dict(visible_text=item.get('old',''),reason='程序测试桩，不代表模型视觉定位',
            source_sha256=obj.get('media_sha256'),upper_objects=[dict(object_id=o['id'],relation='visible_with_target',
                reason='程序测试桩，候选仅相交，不作为实际遮挡证明') for o in ppt_analysis.upper_candidates(objects,obj,clean['page_polygon'])])
    return manifest


def submit_mock_plan(jid,lease,path):
    import json, core, ppt_workflow as ppt
    core.dump(path,add_mock_analysis(json.loads(Path(path).read_text(encoding='utf-8'))))
    return ppt.write_plan(jid,lease,path)


def make_layered_deck(path):
    """Three title cases: covered old title, picture only, unrelated intersection."""
    from pptx import Presentation
    from pptx.util import Inches, Pt
    from pptx.dml.color import RGBColor
    from pptx.enum.shapes import MSO_SHAPE
    p=Presentation();p.slide_width=Inches(12);p.slide_height=Inches(6.75)
    im=Image.new('RGB',(1200,675),'white');draw=ImageDraw.Draw(im)
    draw.text((100,55),'OLD TITLE',font=ImageFont.truetype('C:/Windows/Fonts/arial.ttf',48),fill='#005030')
    data=io.BytesIO();im.save(data,format='PNG')
    for n in range(3):
        s=p.slides.add_slide(p.slide_layouts[6]);s.shapes.add_picture(io.BytesIO(data.getvalue()),0,0,width=p.slide_width,height=p.slide_height)
        if n==0:
            cover=s.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(.7), Inches(.3), Inches(10), Inches(1.1))
            cover.fill.solid();cover.fill.fore_color.rgb=RGBColor.from_string('F8E8D0');cover.line.fill.background()
            text=s.shapes.add_textbox(Inches(.9), Inches(.4), Inches(9), Inches(.9));text.text='NEW TITLE'
            run=text.text_frame.paragraphs[0].runs[0];run.font.size=Pt(36);run.font.color.rgb=RGBColor.from_string('804020')
        if n==2:
            # Intersects picture and proposed title rectangle, but is transparent
            # and only draws a distant edge. Intersection does not hide glyphs.
            shape=s.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(.5), Inches(.2), Inches(11), Inches(3))
            shape.fill.background();shape.line.color.rgb=RGBColor.from_string('A08020')
    p.save(path);return Path(path)
